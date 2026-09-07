import openpyxl
from django.db import transaction

from .models.admin_dashboard_models import Categories, SubCategories, SubSubCategories
from .bulk_import_common import RowError, ImageResolver, _clean, display_row_number

SUB_CATEGORIES_SHEET_NAMES = ["Sub Categories", "Category_Subcategory"]
SUB_SUB_CATEGORIES_SHEET_NAMES = ["Sub Sub Categories"]

TRUTHY_VALUES = {"yes", "y", "true", "1", "x"}


def _find_sheet(wb, candidate_names):
    for name in candidate_names:
        if name in wb.sheetnames:
            return wb[name]
    return None


def _header_indices(header_row, spec):
    cells = [(i, _clean(cell).lower()) for i, cell in enumerate(header_row or [])]
    result = {}
    for field, aliases in spec.items():
        idx = None
        for i, text in cells:
            if text in aliases:
                idx = i
                break
        result[field] = idx
    return result


def _cell(row, idx):
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _clean_or_none(value):
    text = _clean(value)
    return text or None


def _clean_position(value):
    if value is None or _clean(value) == "":
        return None
    try:
        pos = int(value)
    except (TypeError, ValueError):
        return None
    return pos if pos >= 0 else None


def _clean_bool(value, default=False):
    if value is None or _clean(value) == "":
        return default
    return _clean(value).lower() in TRUTHY_VALUES


def _resolve_optional_image(resolver, value, sheet_name, row_num, name_for_message, result):
    try:
        return resolver.resolve(value, "Image", required=False)
    except RowError as e:
        result["image_warnings"].append({
            "sheet": sheet_name,
            "row": row_num,
            "name": name_for_message,
            "message": str(e),
        })
        return None


def run_category_bulk_import(excel_file, images_zip_file=None):
    wb = openpyxl.load_workbook(excel_file, data_only=True)
    resolver = ImageResolver(images_zip_file)

    result = {
        "cat_created": 0,
        "cat_existing": 0,
        "sub_created": 0,
        "sub_existing": 0,
        "sub_skipped": [],
        "subsub_created": 0,
        "subsub_existing": 0,
        "subsub_skipped": [],
        "errors": [],
        "image_warnings": [],
        "sheets_found": list(wb.sheetnames),
        "sheets_missing": [],
    }

    if "Categories" not in wb.sheetnames:
        result["sheets_missing"].append("Categories")
    if not _find_sheet(wb, SUB_CATEGORIES_SHEET_NAMES):
        result["sheets_missing"].append("Sub Categories")
    if not _find_sheet(wb, SUB_SUB_CATEGORIES_SHEET_NAMES):
        result["sheets_missing"].append("Sub Sub Categories")

    # ---------------- Categories ----------------
    if "Categories" in wb.sheetnames:
        ws = wb["Categories"]
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        cols = _header_indices(header_row, {
            "name": ["category name", "category"],
            "description": ["description"],
            "image": ["image"],
            "sl": ["sl", "sl no", "sl no.", "sl.", "serial no", "serial number"],
        })
        name_idx = cols["name"] if cols["name"] is not None else 0

        for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row or all(cell is None for cell in row):
                continue
            display_row = display_row_number(row, cols["sl"], row_num)
            name = _clean(_cell(row, name_idx))
            description = _clean(_cell(row, cols["description"]))
            image_value = _cell(row, cols["image"])
            if not name:
                result["errors"].append({"sheet": "Categories", "row": display_row, "field": "Category Name", "message": "Category Name cannot be left empty"})
                continue

            with transaction.atomic():
                obj, created = Categories.objects.get_or_create(
                    name=name,
                    defaults={"description": description or None},
                )
                if created:
                    image_file = _resolve_optional_image(resolver, image_value, "Categories", display_row, name, result)
                    if image_file:
                        obj.image.save(image_file.name, image_file, save=True)
            if created:
                result["cat_created"] += 1
            else:
                result["cat_existing"] += 1

    # ---------------- Sub Categories ----------------
    ws = _find_sheet(wb, SUB_CATEGORIES_SHEET_NAMES)
    if ws is not None:
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        cols = _header_indices(header_row, {
            "category": ["category name", "category"],
            "sub": ["sub category name", "sub category", "subcategory"],
            "description": ["description"],
            "image": ["image"],
            "column": ["column"],
            "position": ["position"],
            "has_sub_sub_cat": ["has sub sub cat", "has sub sub category"],
            "sl": ["sl", "sl no", "sl no.", "sl.", "serial no", "serial number"],
        })
        cat_idx = cols["category"] if cols["category"] is not None else 0
        sub_idx = cols["sub"] if cols["sub"] is not None else 1

        for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row or all(cell is None for cell in row):
                continue
            display_row = display_row_number(row, cols["sl"], row_num)
            cat_name = _clean(_cell(row, cat_idx))
            sub_name = _clean(_cell(row, sub_idx))
            description = _clean(_cell(row, cols["description"]))
            image_value = _cell(row, cols["image"])
            column_value = _clean_or_none(_cell(row, cols["column"]))
            position_value = _clean_position(_cell(row, cols["position"]))
            has_sub_sub_cat_value = _clean_bool(_cell(row, cols["has_sub_sub_cat"]), default=False)

            if not cat_name or not sub_name:
                result["errors"].append({"sheet": "Sub Categories", "row": display_row, "field": "Category/Sub Category Name", "message": "Both Category Name and Sub Category Name are required"})
                continue

            category = Categories.objects.filter(name=cat_name).first()
            if not category:
                result["sub_skipped"].append({"subcategory": sub_name, "category": cat_name})
                continue

            with transaction.atomic():
                obj, created = SubCategories.objects.get_or_create(
                    categories=category,
                    sub_cat_name=sub_name,
                    defaults={
                        "description": description or None,
                        "column": column_value,
                        "position": position_value,
                        "has_sub_sub_cat": has_sub_sub_cat_value,
                    },
                )
                if created:
                    image_file = _resolve_optional_image(resolver, image_value, "Sub Categories", display_row, sub_name, result)
                    if image_file:
                        obj.image.save(image_file.name, image_file, save=True)
            if created:
                result["sub_created"] += 1
            else:
                result["sub_existing"] += 1

    # ---------------- Sub Sub Categories ----------------
    ws = _find_sheet(wb, SUB_SUB_CATEGORIES_SHEET_NAMES)
    if ws is not None:
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        cols = _header_indices(header_row, {
            "category": ["category name", "category"],
            "sub": ["sub category name", "sub category", "subcategory"],
            "subsub": ["sub sub category name", "sub sub category", "subsubcategory"],
            "image": ["image"],
            "sl": ["sl", "sl no", "sl no.", "sl.", "serial no", "serial number"],
        })
        cat_idx = cols["category"] if cols["category"] is not None else 0
        sub_idx = cols["sub"] if cols["sub"] is not None else 1
        subsub_idx = cols["subsub"] if cols["subsub"] is not None else 2

        for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row or all(cell is None for cell in row):
                continue
            display_row = display_row_number(row, cols["sl"], row_num)
            cat_name = _clean(_cell(row, cat_idx))
            sub_name = _clean(_cell(row, sub_idx))
            subsub_name = _clean(_cell(row, subsub_idx))
            image_value = _cell(row, cols["image"])

            if not cat_name or not sub_name or not subsub_name:
                result["errors"].append({"sheet": "Sub Sub Categories", "row": display_row, "field": "Category/Sub Category/Sub Sub Category Name", "message": "All three names are required"})
                continue

            sub_category = SubCategories.objects.filter(categories__name=cat_name, sub_cat_name=sub_name).first()
            if not sub_category:
                result["subsub_skipped"].append({"sub_sub_category": subsub_name, "category": cat_name, "subcategory": sub_name})
                continue

            with transaction.atomic():
                obj, created = SubSubCategories.objects.get_or_create(
                    sub_categories=sub_category,
                    sub_sub_cat_name=subsub_name,
                )
                if created:
                    image_file = _resolve_optional_image(resolver, image_value, "Sub Sub Categories", display_row, subsub_name, result)
                    if image_file:
                        obj.image.save(image_file.name, image_file, save=True)
            if created:
                result["subsub_created"] += 1
            else:
                result["subsub_existing"] += 1

    result["total_skipped_errors"] = (
        len(result["sub_skipped"]) + len(result["subsub_skipped"]) + len(result["errors"])
    )
    return result