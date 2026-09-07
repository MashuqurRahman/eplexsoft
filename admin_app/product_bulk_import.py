from decimal import Decimal, InvalidOperation

import openpyxl
from django.db import transaction

from .models.admin_dashboard_models import (
    Categories, SubCategories, SubSubCategories, Brand, Unit, Color, Size,
    Product, ProductVarient, ProductAttribute, ProductAttributeImage,
)
from .bulk_import_common import RowError, ImageResolver, _clean, sl_column_index, display_row_number

TRUE_VALUES = {"true", "1", "yes", "y", "on"}

PRODUCTS_HEADER_ALIASES = {
    "product_key":    ["Product Key"],
    "category":       ["Category Name", "Category*", "Category"],
    "sub_category":   ["Sub Category Name", "Sub Category*", "Sub Category", "Subcategory", "Subcategory*"],
    "sub_sub_category": ["Sub Sub Category Name", "Sub Sub Category*", "Sub Sub Category"],
    "product_name":   ["Product Name", "Product Name*"],
    "description":    ["Description", "Description*"],
    "moq":            ["MOQ", "MOQ*"],
    "vat_amount":     ["VAT Amount", "VAT Amount*"],
    "gst_amount":     ["GST Amount", "GST Amount*"],
    "is_applicable":  ["Is VAT/GST Percentage", "Is VAT/GST Percentage (Yes/No)", "Is VAT/GST Percentage*"],
    "is_popular":     ["Is Popular", "Is Popular (Yes/No)", "Is Popular*"],
    "is_active":      ["Is Active", "Active (Yes/No)", "Is Active*"],
    "is_best_deal":   ["Is Best Deal", "Is Best Deal (Yes/No)", "Is Best Deal*"],
    "approval":       ["Approval"],
    "brand":          ["Brand Name", "Brand"],
    "unit":           ["Unit Name", "Unit*", "Unit"],
    "sku":            ["SKU", "SKU*"],
    "cover_photo":    ["Cover Photo", "Cover Photo (filename)*", "Cover Photo (filename)"],
}

VARIANTS_HEADER_ALIASES = {
    "product_key":    ["Product Key"],
    "sku":            ["SKU", "SKU*"],
    "color":          ["Color Name", "Color*", "Color"],
    "size":           ["Size Value", "Size*", "Size"],
    "regular_price":  ["Regular Price", "Regular Price*"],
    "discount_price": ["Discount Price", "Discounted Price", "Discount Price*"],
    "buying_price":   ["Buying Price", "Buying Price*"],
    "height":         ["Height", "Height (in)"],
    "width":          ["Width", "Width (in)"],
    "weight":         ["Weight", "Weight (gm)*", "Weight (gm)"],
    "stock":          ["Stock", "Stock*"],
    "is_cover":       ["Is Cover", "Is Cover (Yes/No)", "Is Cover*"],
    "image":          ["Image", "Product Image (filename)*", "Product Image (filename)"],
    "attribute_images": ["Attribute Images", "Attribute Images (filenames, comma-separated)"],
}


def _to_bool(value, default=False):
    text = _clean(value).lower()
    if not text:
        return default
    return text in TRUE_VALUES


def _to_decimal(value, field_label, required=True, min_value=None):
    text = _clean(value)
    if not text:
        if required:
            raise RowError(f"{field_label} cannot be left empty")
        return None
    try:
        dec = Decimal(text)
    except InvalidOperation:
        raise RowError(f"{field_label} must be a number (got: {text!r})")
    if min_value is not None and dec < min_value:
        raise RowError(f"{field_label} must be at least {min_value}")
    return dec


def _to_float(value):
    text = _clean(value)
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _to_int(value, field_label, default=None, min_value=None):
    text = _clean(value)
    if not text:
        if default is not None:
            return default
        raise RowError(f"{field_label} cannot be left empty")
    try:
        n = int(float(text))
    except ValueError:
        raise RowError(f"{field_label} must be a number (got: {text!r})")
    if min_value is not None and n < min_value:
        raise RowError(f"{field_label} must be at least {min_value}")
    return n


def _build_column_map(ws, alias_groups):
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    normalized = {}
    for idx, val in enumerate(header_row):
        if val is None:
            continue
        normalized[str(val).strip().lower()] = idx

    colmap = {}
    for field, aliases in alias_groups.items():
        for alias in aliases:
            idx = normalized.get(alias.strip().lower())
            if idx is not None:
                colmap[field] = idx
                break
    return colmap


def _get(row, colmap, field):
    idx = colmap.get(field)
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def run_product_bulk_import(excel_file, images_zip_file=None):
    wb = openpyxl.load_workbook(excel_file, data_only=True)
    resolver = ImageResolver(images_zip_file)

    result = {
        "product_created": 0,
        "product_errors": [],
        "variant_created": 0,
        "variant_errors": [],
        "sheets_found": list(wb.sheetnames),
        "sheets_missing": [],
    }
    if "Products" not in wb.sheetnames:
        result["sheets_missing"].append("Products")
    if "Variants" not in wb.sheetnames:
        result["sheets_missing"].append("Variants")

    products_by_key = {}  
    varients_by_key = {}   
    cover_seen = set()     

    # ---------------- Products sheet ----------------
    if "Products" in wb.sheetnames:
        ws = wb["Products"]
        colmap = _build_column_map(ws, PRODUCTS_HEADER_ALIASES)
        uses_sku_as_key = "product_key" not in colmap and "sku" in colmap
        sl_idx = sl_column_index(next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ()))

        for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row or all(c is None for c in row):
                continue
            display_row = display_row_number(row, sl_idx, row_num)

            def g(field):
                return _get(row, colmap, field)

            key = _clean(g("product_key")) or (_clean(g("sku")) if uses_sku_as_key else "")
            try:
                if not key:
                    raise RowError("Product Key (or SKU, if Product Key column is missing) cannot be left empty")
                if key in products_by_key:
                    raise RowError(f"Product Key '{key}' has been used more than once")

                cat_name = _clean(g("category"))
                sub_name = _clean(g("sub_category"))
                subsub_name = _clean(g("sub_sub_category"))
                product_name = _clean(g("product_name"))
                description = _clean(g("description"))
                unit_name = _clean(g("unit"))

                if not cat_name or not sub_name:
                    raise RowError("Both Category Name and Sub Category Name are required")
                if not product_name:
                    raise RowError("Product Name cannot be left empty")
                if not description:
                    raise RowError("Description cannot be left empty")
                if not unit_name:
                    raise RowError("Unit Name cannot be left empty")

                category = Categories.objects.filter(name=cat_name).first()
                if not category:
                    raise RowError(f"Category '{cat_name}' not found — please create it first via Category bulk upload")
                sub_category = SubCategories.objects.filter(categories=category, sub_cat_name=sub_name).first()
                if not sub_category:
                    raise RowError(f"Sub Category '{sub_name}' (under Category '{cat_name}') not found")
                sub_sub_category = None
                if subsub_name:
                    sub_sub_category = SubSubCategories.objects.filter(sub_categories=sub_category, sub_sub_cat_name=subsub_name).first()
                    if not sub_sub_category:
                        raise RowError(f"Sub Sub Category '{subsub_name}' (under Sub Category '{sub_name}') not found")

                moq = _to_int(g("moq"), "MOQ", default=1, min_value=1)
                vat_amount = _to_float(g("vat_amount")) or 0
                gst_amount = _to_float(g("gst_amount")) or 0
                is_applicable = _to_bool(g("is_applicable"), default=False)
                is_popular = _to_bool(g("is_popular"), default=False)
                is_active = _to_bool(g("is_active"), default=True)
                is_best_deal = _to_bool(g("is_best_deal"), default=False)
                approval = _clean(g("approval")).lower() or "not_approved"
                if approval not in ("approved", "not_approved"):
                    raise RowError("Approval column can only contain 'approved' or 'not_approved'")

                brand_name = _clean(g("brand"))
                brand = None
                if brand_name:
                    brand, _ = Brand.objects.get_or_create(name=brand_name)
                unit, _ = Unit.objects.get_or_create(name=unit_name)

                sku = _clean(g("sku")) or None
                if sku and ProductVarient.objects.filter(sku=sku).exists():
                    raise RowError(f"SKU '{sku}' is already in use")

                cover_file = resolver.resolve(g("cover_photo"), "Cover Photo")

                with transaction.atomic():
                    product = Product(
                        categories=category,
                        sub_categories=sub_category,
                        sub_sub_categories=sub_sub_category,
                        product_name=product_name,
                        description=description,
                        moq=moq,
                        vat_tax_amount=vat_amount,
                        gst_amount=gst_amount,
                        approval=approval,
                        is_applicable=is_applicable,
                        is_popular=is_popular,
                        is_active=is_active,
                        is_best_deal=is_best_deal,
                    )
                    product.save()

                    varient = ProductVarient(product=product, brand=brand, unit=unit, sku=sku)
                    varient.cover_photo.save(cover_file.name, cover_file, save=False)
                    varient.save()

                products_by_key[key] = product
                varients_by_key[key] = varient
                result["product_created"] += 1

            except RowError as e:
                result["product_errors"].append({"sheet": "Products", "row": display_row, "field": "-", "message": str(e), "key": key or "-"})
            except Exception as e:
                result["product_errors"].append({"sheet": "Products", "row": display_row, "field": "-", "message": f"Unexpected error: {e}", "key": key or "-"})

    # ---------------- Variants sheet ----------------
    if "Variants" in wb.sheetnames:
        ws = wb["Variants"]
        colmap = _build_column_map(ws, VARIANTS_HEADER_ALIASES)
        uses_sku_as_key = "product_key" not in colmap and "sku" in colmap
        sl_idx = sl_column_index(next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ()))

        for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row or all(c is None for c in row):
                continue
            display_row = display_row_number(row, sl_idx, row_num)

            def g(field):
                return _get(row, colmap, field)

            key = _clean(g("product_key")) or (_clean(g("sku")) if uses_sku_as_key else "")
            try:
                if not key:
                    raise RowError("Product Key (or SKU) cannot be left empty")
                product = products_by_key.get(key)
                varient = varients_by_key.get(key)
                if not product:
                    raise RowError(f"Product Key '{key}' — no successfully created product with this key found in the Products sheet")

                color_name = _clean(g("color"))
                size_value = _clean(g("size"))
                if not color_name or not size_value:
                    raise RowError("Both Color Name and Size Value are required")

                regular_price = _to_decimal(g("regular_price"), "Regular Price", required=True, min_value=Decimal("1"))
                discount_price = _to_decimal(g("discount_price"), "Discount Price", required=False, min_value=Decimal("1"))
                if discount_price is not None and discount_price >= regular_price:
                    raise RowError("Discount Price must be less than Regular Price")
                buying_price = _to_decimal(g("buying_price"), "Buying Price", required=True, min_value=Decimal("1"))
                height = _to_float(g("height"))
                width = _to_float(g("width"))
                weight = _to_float(g("weight"))
                stock = _to_int(g("stock"), "Stock", default=0, min_value=0)
                is_cover = _to_bool(g("is_cover"), default=False)

                image_file = resolver.resolve(g("image"), "Image")

                color, _ = Color.objects.get_or_create(name=color_name)
                size, _ = Size.objects.get_or_create(value=size_value)

                with transaction.atomic():
                    attribute = ProductAttribute(
                        product=product,
                        product_varient=varient,
                        color=color,
                        size=size,
                        regular_price=regular_price,
                        discount_price=discount_price,
                        buying_price=buying_price,
                        height=height,
                        width=width,
                        weight=weight,
                        stock=stock,
                        is_cover=is_cover,
                    )
                    attribute.image.save(image_file.name, image_file, save=False)
                    attribute.save()

                    if is_cover:
                        cover_seen.add(key)

                    extra_images_raw = _clean(g("attribute_images"))
                    if extra_images_raw:
                        for part in extra_images_raw.replace(",", ";").split(";"):
                            part = part.strip()
                            if not part:
                                continue
                            extra_file = resolver.resolve(part, "Attribute Images")
                            attr_image = ProductAttributeImage(product_attribute=attribute)
                            attr_image.image.save(extra_file.name, extra_file, save=False)
                            attr_image.save()

                result["variant_created"] += 1

            except RowError as e:
                result["variant_errors"].append({"sheet": "Variants", "row": display_row, "field": "-", "message": str(e), "key": key or "-"})
            except Exception as e:
                result["variant_errors"].append({"sheet": "Variants", "row": display_row, "field": "-", "message": f"Unexpected error: {e}", "key": key or "-"})

    for key, product in products_by_key.items():
        if key in cover_seen:
            continue
        first_attr = ProductAttribute.objects.filter(product=product).order_by("id").first()
        if first_attr:
            first_attr.is_cover = True
            first_attr.save()

    result["total_errors"] = len(result["product_errors"]) + len(result["variant_errors"])
    return result